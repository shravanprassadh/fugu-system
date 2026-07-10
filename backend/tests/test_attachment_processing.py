"""Format-specific structured attachment extraction tests."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfWriter

from fugu.attachments.processing import process_attachment_bytes


def test_text_and_source_processing_preserve_lines_and_language() -> None:
    text = process_attachment_bytes(extension="txt", content=b"first\nsecond")
    source = process_attachment_bytes(extension="py", content=b"def answer():\n    return 42\n")

    assert text.structured_content == {
        "kind": "text",
        "language": None,
        "line_count": 2,
        "lines": ["first", "second"],
    }
    assert source.structured_content["kind"] == "source_code"
    assert source.structured_content["language"] == "py"
    assert source.structured_content["lines"] == ["def answer():", "    return 42"]


def test_csv_and_xlsx_processing_preserve_table_and_sheet_structure() -> None:
    csv_result = process_attachment_bytes(
        extension="csv",
        content=b"name,score\nAda,98\nLinus,95\n",
    )
    assert csv_result.structured_content["rows"] == [
        ["name", "score"],
        ["Ada", "98"],
        ["Linus", "95"],
    ]

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Results"
    worksheet.append(["name", "score"])
    worksheet.append(["Ada", 98])
    payload = BytesIO()
    workbook.save(payload)
    workbook.close()

    xlsx_result = process_attachment_bytes(extension="xlsx", content=payload.getvalue())
    sheets = xlsx_result.structured_content["sheets"]
    assert isinstance(sheets, list)
    assert sheets[0]["name"] == "Results"
    assert sheets[0]["rows"] == [["name", "score"], ["Ada", 98]]


def test_docx_processing_preserves_paragraphs_and_tables() -> None:
    document = Document()
    document.add_paragraph("Project objective")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Force"
    table.cell(1, 1).text = "125 N"
    payload = BytesIO()
    document.save(payload)

    processed = process_attachment_bytes(extension="docx", content=payload.getvalue())
    assert processed.structured_content["paragraphs"] == ["Project objective"]
    tables = processed.structured_content["tables"]
    assert isinstance(tables, list)
    assert tables[0]["rows"] == [["Item", "Value"], ["Force", "125 N"]]


def test_pdf_processing_preserves_pages_and_reports_empty_text() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    payload = BytesIO()
    writer.write(payload)

    processed = process_attachment_bytes(extension="pdf", content=payload.getvalue())
    assert processed.structured_content["page_count"] == 1
    assert processed.structured_content["pages"] == [{"page_number": 1, "text": ""}]
    assert processed.warnings == ("PDF page 1 contains no extractable text.",)


def test_image_processing_verifies_pixels_and_preserves_dimensions() -> None:
    image = Image.new("RGB", (8, 6))
    payload = BytesIO()
    image.save(payload, format="PNG")

    processed = process_attachment_bytes(extension="png", content=payload.getvalue())
    assert processed.structured_content == {
        "kind": "image",
        "format": "png",
        "width": 8,
        "height": 6,
        "mode": "RGB",
        "native_multimodal": True,
        "text_extracted": False,
    }
    assert processed.warnings == ("Image pixels require a vision-capable Reader model for semantic interpretation.",)
