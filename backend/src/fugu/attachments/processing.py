"""Bounded format-specific attachment extraction preserving document structure."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO, StringIO
from typing import Any

from docx import Document
from openpyxl import load_workbook
from PIL import Image, UnidentifiedImageError
from pypdf import PdfReader

from fugu.storage.errors import AttachmentValidationError

_PROCESSING_VERSION = "document-v1"
_MAX_PDF_PAGES = 200
_MAX_TEXT_LINES = 10_000
_MAX_TEXT_CHARACTERS = 500_000
_MAX_ROWS_PER_TABLE = 2_000
_MAX_COLUMNS_PER_ROW = 100
_MAX_DOCX_PARAGRAPHS = 10_000
_MAX_DOCX_TABLES = 200

_IMAGE_EXTENSIONS = frozenset({"png", "jpg", "jpeg", "webp"})
_SOURCE_EXTENSIONS = frozenset(
    {
        "py",
        "js",
        "jsx",
        "ts",
        "tsx",
        "java",
        "c",
        "h",
        "cpp",
        "hpp",
        "cs",
        "go",
        "rs",
        "rb",
        "php",
        "swift",
        "kt",
        "kts",
        "sql",
        "sh",
        "bash",
        "json",
        "yaml",
        "yml",
        "xml",
        "html",
        "css",
    }
)


@dataclass(frozen=True, slots=True)
class ProcessedAttachment:
    """Structured content and safe warnings ready for relational persistence."""

    structured_content: dict[str, object]
    warnings: tuple[str, ...]
    processing_version: str = _PROCESSING_VERSION


def _bounded_text(value: str, *, warnings: list[str], label: str) -> str:
    if len(value) <= _MAX_TEXT_CHARACTERS:
        return value
    warnings.append(f"{label} was truncated after {_MAX_TEXT_CHARACTERS} characters.")
    return value[:_MAX_TEXT_CHARACTERS]


def _json_value(value: object) -> object:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _process_text(content: bytes, *, extension: str) -> ProcessedAttachment:
    text = content.decode("utf-8-sig")
    warnings: list[str] = []
    lines = text.splitlines()
    if len(lines) > _MAX_TEXT_LINES:
        warnings.append(f"The file was truncated after {_MAX_TEXT_LINES} lines.")
        lines = lines[:_MAX_TEXT_LINES]
    bounded = _bounded_text("\n".join(lines), warnings=warnings, label="Text content")
    return ProcessedAttachment(
        structured_content={
            "kind": "source_code" if extension in _SOURCE_EXTENSIONS else "text",
            "language": extension if extension in _SOURCE_EXTENSIONS else None,
            "line_count": len(lines),
            "lines": bounded.splitlines(),
        },
        warnings=tuple(warnings),
    )


def _process_csv(content: bytes) -> ProcessedAttachment:
    text = content.decode("utf-8-sig")
    reader = csv.reader(StringIO(text))
    warnings: list[str] = []
    rows: list[list[str]] = []
    for row_index, row in enumerate(reader):
        if row_index >= _MAX_ROWS_PER_TABLE:
            warnings.append(f"CSV rows were truncated after {_MAX_ROWS_PER_TABLE} rows.")
            break
        if len(row) > _MAX_COLUMNS_PER_ROW:
            warnings.append(f"CSV row {row_index + 1} was truncated after {_MAX_COLUMNS_PER_ROW} columns.")
        rows.append(row[:_MAX_COLUMNS_PER_ROW])
    return ProcessedAttachment(
        structured_content={
            "kind": "spreadsheet",
            "format": "csv",
            "row_count": len(rows),
            "column_count": max((len(row) for row in rows), default=0),
            "rows": rows,
        },
        warnings=tuple(warnings),
    )


def _process_xlsx(content: bytes) -> ProcessedAttachment:
    warnings: list[str] = []
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheets: list[dict[str, object]] = []
    try:
        for worksheet in workbook.worksheets:
            rows: list[list[object]] = []
            for row_index, row in enumerate(worksheet.iter_rows(values_only=True)):
                if row_index >= _MAX_ROWS_PER_TABLE:
                    warnings.append(f"Sheet {worksheet.title!r} was truncated after {_MAX_ROWS_PER_TABLE} rows.")
                    break
                values = [_json_value(value) for value in row[:_MAX_COLUMNS_PER_ROW]]
                if len(row) > _MAX_COLUMNS_PER_ROW:
                    warnings.append(
                        f"Sheet {worksheet.title!r}, row {row_index + 1} was truncated after "
                        f"{_MAX_COLUMNS_PER_ROW} columns."
                    )
                rows.append(values)
            sheets.append(
                {
                    "name": worksheet.title,
                    "row_count": len(rows),
                    "column_count": max((len(row) for row in rows), default=0),
                    "rows": rows,
                }
            )
    finally:
        workbook.close()
    return ProcessedAttachment(
        structured_content={"kind": "spreadsheet", "format": "xlsx", "sheets": sheets},
        warnings=tuple(warnings),
    )


def _process_docx(content: bytes) -> ProcessedAttachment:
    document = Document(BytesIO(content))
    warnings: list[str] = []
    paragraph_values = [paragraph.text for paragraph in document.paragraphs[:_MAX_DOCX_PARAGRAPHS]]
    if len(document.paragraphs) > _MAX_DOCX_PARAGRAPHS:
        warnings.append(f"DOCX paragraphs were truncated after {_MAX_DOCX_PARAGRAPHS} paragraphs.")
    tables: list[dict[str, object]] = []
    for table_index, table in enumerate(document.tables[:_MAX_DOCX_TABLES]):
        rows: list[list[str]] = []
        for row_index, row in enumerate(table.rows[:_MAX_ROWS_PER_TABLE]):
            values = [cell.text for cell in row.cells[:_MAX_COLUMNS_PER_ROW]]
            if len(row.cells) > _MAX_COLUMNS_PER_ROW:
                warnings.append(
                    f"DOCX table {table_index + 1}, row {row_index + 1} was truncated after "
                    f"{_MAX_COLUMNS_PER_ROW} columns."
                )
            rows.append(values)
        if len(table.rows) > _MAX_ROWS_PER_TABLE:
            warnings.append(f"DOCX table {table_index + 1} was truncated after {_MAX_ROWS_PER_TABLE} rows.")
        tables.append({"index": table_index + 1, "rows": rows})
    if len(document.tables) > _MAX_DOCX_TABLES:
        warnings.append(f"DOCX tables were truncated after {_MAX_DOCX_TABLES} tables.")
    return ProcessedAttachment(
        structured_content={
            "kind": "document",
            "format": "docx",
            "paragraphs": paragraph_values,
            "tables": tables,
        },
        warnings=tuple(warnings),
    )


def _process_pdf(content: bytes) -> ProcessedAttachment:
    reader = PdfReader(BytesIO(content))
    if reader.is_encrypted:
        raise AttachmentValidationError("Encrypted PDFs are not supported.")
    warnings: list[str] = []
    pages: list[dict[str, object]] = []
    for page_index, page in enumerate(reader.pages[:_MAX_PDF_PAGES]):
        text = page.extract_text() or ""
        if not text.strip():
            warnings.append(f"PDF page {page_index + 1} contains no extractable text.")
        pages.append(
            {
                "page_number": page_index + 1,
                "text": _bounded_text(text, warnings=warnings, label=f"PDF page {page_index + 1}"),
            }
        )
    if len(reader.pages) > _MAX_PDF_PAGES:
        warnings.append(f"PDF pages were truncated after {_MAX_PDF_PAGES} pages.")
    metadata: dict[str, object] = {}
    if reader.metadata:
        metadata = {
            str(key).lstrip("/"): _json_value(value) for key, value in reader.metadata.items() if value is not None
        }
    return ProcessedAttachment(
        structured_content={
            "kind": "document",
            "format": "pdf",
            "page_count": len(reader.pages),
            "pages": pages,
            "metadata": metadata,
        },
        warnings=tuple(warnings),
    )


def _process_image(content: bytes) -> ProcessedAttachment:
    try:
        with Image.open(BytesIO(content)) as image:
            image.verify()
        with Image.open(BytesIO(content)) as image:
            width, height = image.size
            image_format = image.format
            mode = image.mode
    except (UnidentifiedImageError, OSError) as exc:
        raise AttachmentValidationError("The image could not be decoded safely.") from exc
    return ProcessedAttachment(
        structured_content={
            "kind": "image",
            "format": (image_format or "unknown").lower(),
            "width": width,
            "height": height,
            "mode": mode,
            "native_multimodal": True,
            "text_extracted": False,
        },
        warnings=("Image pixels require a vision-capable Reader model for semantic interpretation.",),
    )


def process_attachment_bytes(*, extension: str, content: bytes) -> ProcessedAttachment:
    """Extract one accepted attachment into a bounded, format-specific JSON structure."""
    normalized_extension = extension.strip().lower()
    if normalized_extension in _IMAGE_EXTENSIONS:
        return _process_image(content)
    if normalized_extension == "pdf":
        return _process_pdf(content)
    if normalized_extension == "docx":
        return _process_docx(content)
    if normalized_extension == "xlsx":
        return _process_xlsx(content)
    if normalized_extension == "csv":
        return _process_csv(content)
    if normalized_extension in {"txt", "md"} | _SOURCE_EXTENSIONS:
        return _process_text(content, extension=normalized_extension)
    raise AttachmentValidationError(f"No processor is available for .{normalized_extension} files.")
